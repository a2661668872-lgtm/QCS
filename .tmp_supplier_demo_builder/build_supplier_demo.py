from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"D:\工作空间\Obsidian\AI第二大脑\AI第二大脑_内容创作系统")
OUTPUT_DIR = ROOT / "outputs" / "supplier_demo_20260816"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_FILE = OUTPUT_DIR / "供应商名单_虚构演示数据_1200家.xlsx"

ROW_COUNT = 1200

CATEGORIES = [
    ("服饰内衣", "防滑袜"), ("家居日用", "折叠收纳箱"), ("个护清洁", "便携洗脸巾"),
    ("厨房用品", "硅胶保鲜盒"), ("数码配件", "磁吸手机支架"), ("宠物用品", "宠物清洁手套"),
    ("母婴用品", "婴儿防撞条"), ("运动户外", "弹力训练带"), ("办公文具", "桌面收纳架"),
    ("汽车用品", "车载香薰"), ("美容工具", "便携化妆镜"), ("旅行用品", "压缩收纳袋"),
    ("清洁工具", "免手洗拖把"), ("小家电", "迷你加湿器"), ("礼品包装", "节日礼盒"),
    ("五金工具", "多功能螺丝刀"), ("鞋靴配件", "防磨脚贴"), ("食品包装", "密封包装袋"),
    ("户外照明", "便携露营灯"), ("健康用品", "热敷眼罩"),
]
REGIONS = [
    "浙江义乌", "浙江宁波", "浙江温州", "浙江台州", "广东深圳", "广东广州", "广东东莞", "广东佛山",
    "福建泉州", "福建厦门", "江苏苏州", "江苏南通", "山东临沂", "河北白沟", "河南郑州", "江西南昌",
]
PLATFORMS = ["1688", "义乌购", "中国制造网", "慧聪网", "企业官网"]
SURNAMES = ["陈", "林", "王", "李", "张", "刘", "黄", "赵", "周", "吴"]

HEADERS = [
    "编号", "供应商名称", "主营品类", "代表产品", "地区", "来源平台", "联系人", "公开电话（演示）", "微信（演示）",
    "最低价（元）", "最高价（元）", "起订量", "支持拿样", "支持定制", "邮箱（演示）", "来源链接（演示）",
    "优先级", "综合评分", "推荐理由", "跟进状态", "数据属性",
]

TEAL = "0F766E"
TEAL_DARK = "134E4A"
TEAL_PALE = "F0FDFA"
AMBER_PALE = "FEF3C7"
AMBER_TEXT = "92400E"
WHITE = "FFFFFF"
TEXT = "1F2937"
LIGHT_LINE = "D1FAE5"


def build_raw_row(i: int) -> list:
    category, product = CATEGORIES[(i - 1) % len(CATEGORIES)]
    region = REGIONS[(i * 3 + 1) % len(REGIONS)]
    platform = PLATFORMS[(i * 7 + 2) % len(PLATFORMS)]
    min_price = round(2.8 + ((i * 17) % 460) / 10, 1)
    max_price = round(min_price + 4 + ((i * 11) % 180) / 10, 1)
    moq = [20, 50, 100, 200, 300, 500][(i * 5) % 6]
    sample = "否" if i % 4 == 0 else "是"
    custom = "否" if i % 3 == 0 else "是"
    number = f"{i:04d}"
    city = region[2:]
    return [
        f"GYS-{number}",
        f"{city}{category}供应链{number}有限公司（演示）",
        category,
        product,
        region,
        platform,
        f"{SURNAMES[i % len(SURNAMES)]}经理",
        f"000-0000-{number}",
        f"DEMO_GYS_{number}",
        min_price,
        max_price,
        moq,
        sample,
        custom,
        f"supplier{number}@example.com",
        f"https://example.com/supplier/{number}",
        None,
        None,
        None,
        "未联系",
        "虚构演示数据",
    ]


def compute_score(row: list) -> int:
    platform, min_price, max_price, moq, sample, custom = row[5], row[9], row[10], row[11], row[12], row[13]
    score = 50
    score += 10 if sample == "是" else 0
    score += 10 if custom == "是" else 0
    score += 10 if moq <= 100 else (5 if moq <= 300 else 0)
    score += 10 if min_price <= 20 else (5 if min_price <= 50 else 0)
    score += 5 if max_price - min_price <= 20 else 0
    score += 5 if platform == "企业官网" else 0
    return min(score, 100)


def computed_fields(row: list) -> tuple[str, int, str]:
    score = compute_score(row)
    priority = "A" if score >= 85 else ("B" if score >= 70 else "C")
    reason = (
        "价格与起订量有优势，建议优先询盘" if priority == "A" else
        "条件较均衡，可作为备选" if priority == "B" else
        "起订量或价格需要进一步确认"
    )
    return priority, score, reason


raw_rows = [build_raw_row(i) for i in range(1, ROW_COUNT + 1)]
computed_rows = []
for row in raw_rows:
    priority, score, reason = computed_fields(row)
    view_row = row.copy()
    view_row[16], view_row[17], view_row[18] = priority, score, reason
    computed_rows.append(view_row)

wb = Workbook()
ws_summary = wb.active
ws_summary.title = "概览"
ws = wb.create_sheet("供应商清单")
wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"

thin_light = Side(style="thin", color=LIGHT_LINE)

# 供应商清单
ws.sheet_view.showGridLines = False
ws.merge_cells("A1:U1")
ws["A1"] = "短视频带货供应商调研表｜1,200家演示数据"
ws.merge_cells("A2:U2")
ws["A2"] = "⚠ 本表仅用于录屏与剪辑展示；商家、联系方式和链接均为虚构演示数据，不可用于真实采购或联络。"

for col, header in enumerate(HEADERS, 1):
    ws.cell(4, col, header)

for excel_row, raw in enumerate(raw_rows, 5):
    for col, value in enumerate(raw, 1):
        if col not in (17, 18, 19):
            ws.cell(excel_row, col, value)
    ws.cell(excel_row, 18, f'=MIN(100,50+IF(M{excel_row}="是",10,0)+IF(N{excel_row}="是",10,0)+IF(L{excel_row}<=100,10,IF(L{excel_row}<=300,5,0))+IF(J{excel_row}<=20,10,IF(J{excel_row}<=50,5,0))+IF(K{excel_row}-J{excel_row}<=20,5,0)+IF(F{excel_row}="企业官网",5,0))')
    ws.cell(excel_row, 17, f'=IF(R{excel_row}>=85,"A",IF(R{excel_row}>=70,"B","C"))')
    ws.cell(excel_row, 19, f'=IF(Q{excel_row}="A","价格与起订量有优势，建议优先询盘",IF(Q{excel_row}="B","条件较均衡，可作为备选","起订量或价格需要进一步确认"))')

ws["A1"].fill = PatternFill("solid", fgColor=TEAL)
ws["A1"].font = Font(name="Microsoft YaHei", size=18, bold=True, color=WHITE)
ws["A1"].alignment = Alignment(vertical="center", horizontal="left")
ws["A2"].fill = PatternFill("solid", fgColor=AMBER_PALE)
ws["A2"].font = Font(name="Microsoft YaHei", size=10, italic=True, color=AMBER_TEXT)
ws["A2"].alignment = Alignment(vertical="center", horizontal="left")

for cell in ws[4]:
    cell.fill = PatternFill("solid", fgColor=TEAL_DARK)
    cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=WHITE)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for row in ws.iter_rows(min_row=5, max_row=1204, min_col=1, max_col=21):
    for cell in row:
        cell.font = Font(name="Microsoft YaHei", size=10, color=TEXT)
        cell.alignment = Alignment(vertical="center")

for row in range(5, 1205):
    ws.cell(row, 10).number_format = '¥0.0'
    ws.cell(row, 11).number_format = '¥0.0'
    ws.cell(row, 12).number_format = '#,##0'
    ws.cell(row, 18).number_format = '0'

ws.row_dimensions[1].height = 30
ws.row_dimensions[2].height = 23
ws.row_dimensions[4].height = 30
for row in range(5, 1205):
    ws.row_dimensions[row].height = 21

widths = {
    "A": 12, "B": 34, "C": 14, "D": 18, "E": 13, "F": 12, "G": 10, "H": 18, "I": 18,
    "J": 12, "K": 12, "L": 10, "M": 10, "N": 10, "O": 28, "P": 36, "Q": 9, "R": 10,
    "S": 34, "T": 12, "U": 16,
}
for col, width in widths.items():
    ws.column_dimensions[col].width = width

ws.freeze_panes = "C5"
table = Table(displayName="SupplierDemoTable", ref="A4:U1204")
table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
ws.add_table(table)

dv = DataValidation(type="list", formula1='"未联系,已联系,待复核"', allow_blank=False)
ws.add_data_validation(dv)
dv.add("T5:T1204")

ws.conditional_formatting.add("Q5:Q1204", FormulaRule(formula=['Q5="A"'], fill=PatternFill("solid", fgColor="DCFCE7"), font=Font(bold=True, color="166534")))
ws.conditional_formatting.add("Q5:Q1204", FormulaRule(formula=['Q5="B"'], fill=PatternFill("solid", fgColor="FEF3C7"), font=Font(bold=True, color="92400E")))
ws.conditional_formatting.add("Q5:Q1204", FormulaRule(formula=['Q5="C"'], fill=PatternFill("solid", fgColor="FEE2E2"), font=Font(bold=True, color="991B1B")))

# 概览
ws_summary.sheet_view.showGridLines = False
ws_summary.merge_cells("A1:H1")
ws_summary["A1"] = "供应商调研结果概览（演示）"
ws_summary.merge_cells("A2:H2")
ws_summary["A2"] = "用于视频画面展示｜详细数据见“供应商清单”｜全部联系方式均为虚构占位"

for col, value in enumerate(["总供应商", None, "A级优先", None, "支持拿样", None, "支持定制", None], 1):
    ws_summary.cell(4, col, value)
for cell, formula in {
    "A5": "=COUNTA('供应商清单'!$A$5:$A$1204)",
    "C5": '=COUNTIF(\'供应商清单\'!$Q$5:$Q$1204,"A")',
    "E5": '=COUNTIF(\'供应商清单\'!$M$5:$M$1204,"是")',
    "G5": '=COUNTIF(\'供应商清单\'!$N$5:$N$1204,"是")',
}.items():
    ws_summary[cell] = formula

ws_summary["A7"], ws_summary["B7"] = "来源平台", "供应商数量"
for row, platform in enumerate(PLATFORMS, 8):
    ws_summary.cell(row, 1, platform)
    ws_summary.cell(row, 2, f'=COUNTIF(\'供应商清单\'!$F$5:$F$1204,A{row})')

ws_summary["D7"], ws_summary["E7"] = "主营品类", "供应商数量"
for row, (category, _) in enumerate(CATEGORIES[:8], 8):
    ws_summary.cell(row, 4, category)
    ws_summary.cell(row, 5, f'=COUNTIF(\'供应商清单\'!$C$5:$C$1204,D{row})')

ws_summary["G7"], ws_summary["H7"] = "使用说明", "内容"
notes = [
    ("数据性质", "虚构演示"),
    ("适用场景", "录屏、剪辑、界面展示"),
    ("不可用于", "真实采购与商家联络"),
    ("建议画面", "先概览，再快速滚动清单"),
    ("隐私处理", "可由创作者后期统一打码"),
]
for row, (label, value) in enumerate(notes, 8):
    ws_summary.cell(row, 7, label)
    ws_summary.cell(row, 8, value)

ws_summary["A1"].fill = PatternFill("solid", fgColor=TEAL)
ws_summary["A1"].font = Font(name="Microsoft YaHei", size=20, bold=True, color=WHITE)
ws_summary["A1"].alignment = Alignment(vertical="center")
ws_summary["A2"].fill = PatternFill("solid", fgColor="CCFBF1")
ws_summary["A2"].font = Font(name="Microsoft YaHei", size=10, color="115E59")
ws_summary["A2"].alignment = Alignment(vertical="center")

for cell in ws_summary[4]:
    cell.fill = PatternFill("solid", fgColor=TEAL_DARK)
    cell.font = Font(name="Microsoft YaHei", bold=True, color=WHITE)
    cell.alignment = Alignment(horizontal="center", vertical="center")
for cell in ws_summary[5]:
    cell.fill = PatternFill("solid", fgColor=TEAL_PALE)
    cell.font = Font(name="Microsoft YaHei", size=18, bold=True, color=TEAL)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = Border(top=thin_light, bottom=thin_light, left=thin_light, right=thin_light)

for header_range in ("A7:B7", "D7:E7", "G7:H7"):
    for row in ws_summary[header_range]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=TEAL)
            cell.font = Font(name="Microsoft YaHei", bold=True, color=WHITE)
            cell.alignment = Alignment(horizontal="center", vertical="center")

for range_ref in ("A8:B12", "D8:E15", "G8:H12"):
    for row in ws_summary[range_ref]:
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10, color=TEXT)
            cell.border = Border(bottom=thin_light)
            cell.alignment = Alignment(vertical="center")

for col, width in {"A": 16, "B": 14, "C": 16, "D": 14, "E": 16, "F": 14, "G": 16, "H": 30}.items():
    ws_summary.column_dimensions[col].width = width
ws_summary.row_dimensions[1].height = 34
ws_summary.row_dimensions[2].height = 22
ws_summary.row_dimensions[4].height = 22
ws_summary.row_dimensions[5].height = 34
ws_summary.freeze_panes = "A3"

wb.save(OUTPUT_FILE)

# 文件与结构验证
with ZipFile(OUTPUT_FILE, "r") as archive:
    bad_member = archive.testzip()
    if bad_member is not None:
        raise RuntimeError(f"损坏的xlsx成员: {bad_member}")

check_wb = load_workbook(OUTPUT_FILE, data_only=False, read_only=False)
check_ws = check_wb["供应商清单"]
if check_ws.max_row != 1204 or check_ws.max_column != 21:
    raise RuntimeError(f"清单尺寸异常: {check_ws.max_row}x{check_ws.max_column}")
if check_ws["A1204"].value != "GYS-1200":
    raise RuntimeError("末行编号校验失败")
if not str(check_ws["Q5"].value).startswith("=IF("):
    raise RuntimeError("优先级公式缺失")
if not str(check_ws["R1204"].value).startswith("=MIN("):
    raise RuntimeError("评分公式未填充到末行")

error_tokens = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")
formula_error_hits = []
for sheet in check_wb.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and any(token in cell.value for token in error_tokens):
                formula_error_hits.append(f"{sheet.title}!{cell.coordinate}")
if formula_error_hits:
    raise RuntimeError(f"发现公式错误标记: {formula_error_hits[:10]}")

# 生成两张紧凑预览图，供视觉检查。
font_path = Path(r"C:\Windows\Fonts\msyh.ttc")
font_bold_path = Path(r"C:\Windows\Fonts\msyhbd.ttc")
font = ImageFont.truetype(str(font_path), 18)
font_small = ImageFont.truetype(str(font_path), 15)
font_bold = ImageFont.truetype(str(font_bold_path if font_bold_path.exists() else font_path), 18)
font_title = ImageFont.truetype(str(font_bold_path if font_bold_path.exists() else font_path), 28)


def draw_table_preview(path: Path, title: str, headers: list[str], data: list[list], widths_px: list[int], max_rows: int) -> None:
    row_h, title_h = 38, 64
    width = sum(widths_px)
    height = title_h + row_h * (max_rows + 1)
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, width, title_h), fill="#0F766E")
    draw.text((18, 14), title, font=font_title, fill="white")
    x = 0
    for idx, header in enumerate(headers):
        w = widths_px[idx]
        draw.rectangle((x, title_h, x + w, title_h + row_h), fill="#134E4A", outline="#D1FAE5")
        draw.text((x + 7, title_h + 8), str(header)[:18], font=font_bold, fill="white")
        x += w
    for r_idx, row in enumerate(data[:max_rows]):
        y = title_h + row_h * (r_idx + 1)
        fill = "#F0FDFA" if r_idx % 2 == 0 else "#FFFFFF"
        x = 0
        for c_idx, value in enumerate(row):
            w = widths_px[c_idx]
            draw.rectangle((x, y, x + w, y + row_h), fill=fill, outline="#D1FAE5")
            text = "" if value is None else str(value)
            max_chars = max(5, int(w / 16))
            if len(text) > max_chars:
                text = text[: max_chars - 1] + "…"
            draw.text((x + 6, y + 9), text, font=font_small, fill="#1F2937")
            x += w
    img.save(path)


summary_counts = {
    "总供应商": ROW_COUNT,
    "A级优先": sum(1 for row in computed_rows if row[16] == "A"),
    "支持拿样": sum(1 for row in computed_rows if row[12] == "是"),
    "支持定制": sum(1 for row in computed_rows if row[13] == "是"),
}
summary_preview_data = [[k, v] for k, v in summary_counts.items()]
draw_table_preview(
    OUTPUT_DIR / "preview_summary.png",
    "供应商调研结果概览（演示）",
    ["指标", "数量"],
    summary_preview_data,
    [360, 220],
    len(summary_preview_data),
)

preview_columns = [0, 1, 2, 3, 4, 5, 6, 7, 8, 16, 17]
preview_headers = [HEADERS[i] for i in preview_columns]
preview_rows = [[row[i] for i in preview_columns] for row in computed_rows[:12]]
draw_table_preview(
    OUTPUT_DIR / "preview_list.png",
    "短视频带货供应商清单｜前12行",
    preview_headers,
    preview_rows,
    [120, 330, 130, 160, 130, 130, 100, 180, 180, 90, 110],
    12,
)

print(f"OUTPUT={OUTPUT_FILE}")
print(f"ROWS={ROW_COUNT}")
print(f"A_PRIORITY={summary_counts['A级优先']}")
print(f"SAMPLE_YES={summary_counts['支持拿样']}")
print(f"CUSTOM_YES={summary_counts['支持定制']}")
print("FORMULA_ERROR_HITS=0")
print("ZIP_TEST=PASS")
